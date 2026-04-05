import streamlit as st
import docx
import json
import re
import pandas as pd
import io
import uuid
import base64
import zipfile
from PIL import Image

# ==========================================
# 內部輔助函數 (文字、圖片壓縮與標籤清洗)
# ==========================================
def compress_image_blob(blob):
    """智慧圖片壓縮引擎：自動縮圖並轉為高壓縮 JPEG"""
    try:
        img = Image.open(io.BytesIO(blob))
        if img.mode != "RGB":
            img = img.convert("RGB")
        img.thumbnail((800, 800), Image.Resampling.LANCZOS)
        out = io.BytesIO()
        img.save(out, format="JPEG", quality=75) 
        return out.getvalue()
    except Exception as e:
        return blob

def get_para_text_with_images(para, image_db):
    """提取文字，並留下 [IMG: 檔名.jpg] 的輕量化佔位符"""
    full_text = para.text  
    img_placeholders = ""
    for run in para.runs:
        try:
            blips = run._element.xpath('.//*[local-name()="blip"]')
            imagedatas = run._element.xpath('.//*[local-name()="imagedata"]')
            
            for element in blips + imagedatas:
                rId = element.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed') or \
                      element.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                if rId:
                    part = para.part.related_parts[rId]
                    compressed_blob = compress_image_blob(part.blob)
                    
                    img_filename = f"IMG_{uuid.uuid4().hex[:8]}.jpg"
                    image_db[img_filename] = compressed_blob
                    img_placeholders += f"\n[IMG: {img_filename}]\n"
        except:
            pass
    return full_text + img_placeholders

def normalize_line(text):
    text = text.replace('（', '(').replace('）', ')')
    text = text.replace('：', ':')
    text = re.sub(r'[ \t\u3000]+', ' ', text)
    return text.strip()

def extract_and_remove_tags(text, q_dict):
    """提取標籤，還給本文乾淨的空間"""
    if not text: return text
    keywords = ["難度", "再現性", "主題", "分類", "單元", "章節"]
    for kw in keywords:
        pattern = r'[\"\'”’]?\s*(' + kw + r')\s*[:：]\s*([^\"\'”’\n]+)[\"\'”’]?'
        for m in re.finditer(pattern, text):
            val = m.group(2).strip()
            val = re.sub(r'\(.*?\)|（.*?）', '', val).strip()
            q_dict["tags"][kw] = val.rstrip(',，').strip()
            text = text.replace(m.group(0), '')
    text = re.sub(r'^[,\"\'\s，]+|[,\"\'\s，]+$', '', text)
    return re.sub(r',\s*,', ',', text).strip()

def _extract_tags_from_all(q_dict):
    for k in ["分類", "單元", "章節"]:
        if k in q_dict["tags"] and "主題" not in q_dict["tags"]:
            q_dict["tags"]["主題"] = q_dict["tags"][k]

def auto_categorize(q_dict, mapping):
    """
    優化後的分類邏輯：
    1. 絕對優先比對『題目本文』。
    2. 只有在題目完全抓不到關鍵字時，才比對『解析的前 50 個字』。
    3. 排除解析後段的干擾。
    """
    if q_dict["tags"].get("主題", "未分類") != "未分類": 
        return 

    # 取得純淨的題目與解析（移除圖片佔位符以利比對）
    q_text = re.sub(r'\[IMG: [^\]]+\]', '', q_dict.get("question_text", "")).lower()
    exp_text = re.sub(r'\[IMG: [^\]]+\]', '', q_dict.get("explanation", "")).lower()
    
    # 步驟 1: 僅針對題目比對
    for topic, keywords in mapping.items():
        if any(kw.lower() in q_text for kw in keywords):
            q_dict["tags"]["主題"] = topic
            return 

    # 步驟 2: 若題目沒抓到，才看解析的最前面（通常解析開頭會重複臨床診斷名稱）
    short_exp = exp_text[:50] 
    for topic, keywords in mapping.items():
        if any(kw.lower() in short_exp for kw in keywords):
            q_dict["tags"]["主題"] = topic
            return

def finalize_question(q_dict, topic_mapping):
    """🌟 終極大水桶分割器：智慧對齊圖片與解析"""
    raw = q_dict.pop("_raw_text", "")
    
    # 1. 優先將標籤洗掉，保持文字純淨
    raw = extract_and_remove_tags(raw, q_dict)
    
    # 2. 尋找選項 A 的起點
    match_A = re.search(r'\(\s*[A]\s*\)(?=.*?\(\s*[B]\s*\))', raw, re.DOTALL)
    if match_A:
        q_dict["question_text"] = raw[:match_A.start()].strip()
        rest_after_A = raw[match_A.start():]
        
        # 3. 尋找「解析」的起點
        exp_pattern = re.compile(r'([\"\'\,\.\-\s【\[<]*)(解\s*答?\s*析)(?!度)[\s:：\]】>\"\',，]*(.*)', re.IGNORECASE | re.DOTALL)
        exp_match = exp_pattern.search(rest_after_A)
        
        if exp_match:
            options_part = rest_after_A[:exp_match.start()].strip()
            explanation_part = rest_after_A[exp_match.start(2):].strip()
            
            # 🌟 核心防漏機制：檢查選項區塊的「最後面」是否有圖片？
            # 如果有，強制把它移動到「解析」的最前方！
            trailing_match = re.search(r'((?:\[IMG:\s*[^\]]+\]\s*)+)$', options_part)
            if trailing_match:
                explanation_part = trailing_match.group(1).strip() + "\n\n" + explanation_part
                options_part = options_part[:trailing_match.start()].strip()
                
            q_dict["explanation"] = explanation_part
        else:
            # 如果整題找不到「解析」兩個字，但選項最後面掛著圖片，也把它當作解析！
            trailing_match = re.search(r'((?:\[IMG:\s*[^\]]+\]\s*)+)$', rest_after_A)
            if trailing_match:
                q_dict["explanation"] = trailing_match.group(1).strip()
                options_part = rest_after_A[:trailing_match.start()].strip()
            else:
                q_dict["explanation"] = ""
                options_part = rest_after_A
                
        # 4. 抽取選項 A, B, C, D
        opt_pattern = re.compile(r'\(\s*(?P<key>[A-E])\s*\)\s*(?P<val>.*?)(?=(?:\(\s*[A-E]\s*\))|$)', re.DOTALL)
        options = {}
        for m in opt_pattern.finditer(options_part):
            options[m.group('key')] = m.group('val').replace('\n', ' ').strip()
        q_dict["options"] = options
        
    else:
        # 如果根本沒有選項
        exp_pattern = re.compile(r'([\"\'\,\.\-\s【\[<]*)(解\s*答?\s*析)(?!度)[\s:：\]】>\"\',，]*(.*)', re.IGNORECASE | re.DOTALL)
        exp_match = exp_pattern.search(raw)
        if exp_match:
            q_part = raw[:exp_match.start()].strip()
            exp_part = raw[exp_match.start(2):].strip()
            
            trailing_match = re.search(r'((?:\[IMG:\s*[^\]]+\]\s*)+)$', q_part)
            if trailing_match:
                exp_part = trailing_match.group(1).strip() + "\n\n" + exp_part
                q_part = q_part[:trailing_match.start()].strip()
                
            q_dict["question_text"] = q_part
            q_dict["explanation"] = exp_part
        else:
            trailing_match = re.search(r'((?:\[IMG:\s*[^\]]+\]\s*)+)$', raw)
            if trailing_match:
                q_dict["explanation"] = trailing_match.group(1).strip()
                q_dict["question_text"] = raw[:trailing_match.start()].strip()
            else:
                q_dict["question_text"] = raw.strip()
                q_dict["explanation"] = ""
        q_dict["options"] = {}

    _extract_tags_from_all(q_dict)
    auto_categorize(q_dict, topic_mapping)

def create_zip_package(json_data, image_db, filename_prefix):
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "a", zipfile.ZIP_DEFLATED, False) as zip_file:
        json_str = json.dumps(json_data, ensure_ascii=False, separators=(',', ':'))
        zip_file.writestr(f"{filename_prefix}.json", json_str.encode('utf-8'))
        for img_name, img_bytes in image_db.items():
            zip_file.writestr(f"images/{img_name}", img_bytes)
    return zip_buffer.getvalue()

# ==========================================
# 網頁介面開始
# ==========================================
st.set_page_config(page_title="國考題庫轉檔系統", page_icon="⚙️", layout="wide")
st.title("⚙️ 國考題庫轉檔系統 (V20 智慧圖文對齊版)")
st.write("已修正圖文分離錯誤！智慧掃描能完美判斷圖片歸屬，讓解析圖文並茂。")

tab1, tab2, tab3 = st.tabs(["🚀 一鍵產出 ZIP 題庫包", "📝 階段一：轉為 Excel 供校對", "💾 階段二：Excel 打包 ZIP"])

default_mapping = {
        "先天免疫與發炎反應": ["發炎", "白血球", "吞噬", "巨噬細胞", "MACROPHAGE", "NEUTROPHIL", "NK細胞", "自然殺手", "TLR", "先天免疫", "發燒", "C-REACTIVE", "CRP"],
        "補體系統": ["補體", "COMPLEMENT", "C3", "C4", "C5", "MAC", "古典途徑", "替代途徑", "凝集素途徑", "C1Q"],
        "抗體與免疫球蛋白": ["抗體", "免疫球蛋白", "IGG", "IGA", "IGM", "IGE", "IGD", "ISOTYPE", "輕鏈", "重鏈", "FAB", "FC", "ALLOTYPE", "IDIOTYPE"],
        "T細胞與細胞免疫": ["T細胞", "T CELL", "CD4", "CD8", "TH1", "TH2", "TREG", "胸腺", "細胞激素", "CYTOKINE", "IL-", "干擾素", "IFN", "穿孔素", "FAS"],
        "B細胞與體液免疫": ["B細胞", "B CELL", "漿細胞", "PLASMA CELL", "記憶B", "BCR"],
        "MHC與移植免疫": ["MHC", "HLA", "移植", "排斥", "組織相容", "GRAFT", "GVHD"],
        "過敏反應": ["過敏", "HYPERSENSITIVITY", "氣喘", "肥大細胞", "MAST CELL", "組織胺", "第一型", "第二型", "第三型", "第四型", "ARTHUS", "接觸性皮膚炎"],
        "自體免疫疾病": ["自體免疫", "AUTOIMMUNE", "ANA", "SLE", "紅斑性狼瘡", "類風濕", "RF", "重症肌無力", "橋本氏", "HASHIMOTO", "GRAVES", "SCLERODERMA", "硬皮症", "乾燥症"],
        "腫瘤免疫": ["腫瘤", "癌症", "TUMOR", "CANCER", "癌", "CEA", "AFP", "PSA", "腫瘤標記", "免疫檢查點", "PD-1", "CTLA-4", "CAR-T", "SIPULEUCEL"],
        "疫苗與預防接種": ["疫苗", "VACCINE", "佐劑", "ADJUVANT", "被動免疫", "主動免疫", "減毒", "類毒素", "TOXOID"],
        "免疫檢驗技術": ["ELISA", "流式細胞", "FLOW CYTOMETRY", "螢光", "沉澱", "凝集", "西方墨點", "WESTERN BLOT", "免疫分析", "RIA", "VDRL", "RPR", "免疫電泳"]

}

def parse_word_document(uploaded_file, topic_mapping):
    doc = docx.Document(uploaded_file)
    image_db = {}
    all_lines = []
    
    # 穿透表格，扁平化所有文字與圖片
    for block in doc.element.body.iterchildren():
        if block.tag.endswith('p'):
            para = docx.text.paragraph.Paragraph(block, doc)
            raw_text = get_para_text_with_images(para, image_db)
            for line in re.split(r'[\n\v]', raw_text):
                clean_line = normalize_line(line)
                if clean_line: all_lines.append(clean_line)
                
        elif block.tag.endswith('tbl'):
            table = docx.table.Table(block, doc)
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        raw_text = get_para_text_with_images(para, image_db)
                        for line in re.split(r'[\n\v]', raw_text):
                            clean_line = normalize_line(line)
                            if clean_line: all_lines.append(clean_line)

    questions = []
    current_year = "未知年份"
    current_topic = "未分類"
    current_q = None

    year_pattern = re.compile(r'(\d{2,4})\s*年')
    q_start_pattern = re.compile(r'^.*?[\(]\s*(?P<ans>[A-Ea-e,皆全對送分]+)\s*[\)]\s*(?P<num>\d+)\s*[.、\s]\s*(?P<text>.*)')
    topic_pattern = re.compile(r'^(?:【([^】]+)】|(?:\w{2}[:：]\s*)(.+))$')

    for text in all_lines:
        t_match = topic_pattern.match(text)
        if t_match and not q_start_pattern.search(text):
            if current_q:
                finalize_question(current_q, topic_mapping)
                questions.append(current_q)
                current_q = None
            current_topic = t_match.group(1) or t_match.group(2)
            continue

        year_match = year_pattern.search(text)
        if year_match and not q_start_pattern.search(text): 
            if current_q:
                finalize_question(current_q, topic_mapping)
                questions.append(current_q)
                current_q = None
            current_year = text.replace('"', '').replace(',', '').strip()
            continue
            
        q_match = q_start_pattern.match(text)
        if q_match:
            if current_q:
                finalize_question(current_q, topic_mapping)
                questions.append(current_q)
            ans = q_match.group('ans').strip().upper().replace('，', ',')
            q_num = int(q_match.group('num')) if q_match.group('num').isdigit() else 0
            current_q = {
                "question_number": q_num, "answer": ans, "explanation": "", 
                "tags": {"年份": current_year, "主題": current_topic}, 
                "_raw_text": q_match.group('text').strip()
            }
            continue
            
        if current_q:
            current_q["_raw_text"] += "\n" + text

    if current_q:
        finalize_question(current_q, topic_mapping)
        questions.append(current_q)
        
    return questions, image_db

with tab1:
    st.info("直接將 Word 轉換為最終 ZIP 題庫包。完美校正圖片位置，讓解析完整顯示。")
    mapping_str_1 = st.text_area("關鍵字分類字典：", value=json.dumps(default_mapping, ensure_ascii=False, indent=4), height=150, key="map1")
    try: topic_mapping_1 = json.loads(mapping_str_1)
    except: topic_mapping_1 = default_mapping
    
    uploaded_word_1 = st.file_uploader("上傳 Word 題庫 (.docx)", type=["docx"], key="w1")
    if uploaded_word_1 and st.button("🚀 產出 ZIP 題庫包", type="primary", use_container_width=True):
        with st.spinner("正在執行智慧對齊與打包作業..."):
            qs, img_db = parse_word_document(uploaded_word_1, topic_mapping_1)
            if qs:
                prefix = uploaded_word_1.name.replace(".docx", "")
                zip_bytes = create_zip_package(qs, img_db, prefix)
                st.success(f"成功解析 {len(qs)} 題！共抽取 {len(img_db)} 張圖片，已完美分離打包！")
                st.download_button("📦 下載 ZIP 題庫包", data=zip_bytes, file_name=f"{prefix}_題庫包.zip", mime="application/zip", type="primary", use_container_width=True)

with tab2:
    st.info("讓老師用 Excel 校對分類。⚠️ 圖片會以 [IMG_xxx.jpg] 顯示在解析欄位中，轉回 JSON 時會自動復原！")
    mapping_str_2 = st.text_area("關鍵字分類字典：", value=json.dumps(default_mapping, ensure_ascii=False, indent=4), height=150, key="map2")
    try: topic_mapping_2 = json.loads(mapping_str_2)
    except: topic_mapping_2 = default_mapping
    
    uploaded_word_2 = st.file_uploader("上傳 Word 題庫 (.docx)", type=["docx"], key="w2")
    if uploaded_word_2 and st.button("🚀 產出待校對 Excel 與圖片暫存檔", type="primary", use_container_width=True):
        with st.spinner("正在產生 Excel 並暫存壓縮圖片..."):
            qs, img_db = parse_word_document(uploaded_word_2, topic_mapping_2)
            if qs:
                excel_rows = []
                for q in qs:
                    opts = q.get("options", {})
                    excel_rows.append({
                        "年份": q["tags"].get("年份", ""), "題號": q.get("question_number", ""),
                        "主題 (下拉選單)": q["tags"].get("主題", ""), "題目": q.get("question_text", ""),
                        "選項A": opts.get("A", ""), "選項B": opts.get("B", ""),
                        "選項C": opts.get("C", ""), "選項D": opts.get("D", ""),
                        "正確答案": q.get("answer", ""), "解析": q.get("explanation", "")
                    })
                df = pd.DataFrame(excel_rows)
                output = io.BytesIO()
                with pd.ExcelWriter(output, engine='xlsxwriter') as writer:
                    df.to_excel(writer, index=False, sheet_name='待校對題庫')
                    workbook = writer.book
                    worksheet = writer.sheets['待校對題庫']
                    worksheet.set_column('A:B', 8); worksheet.set_column('C:C', 20); worksheet.set_column('D:D', 45)
                    worksheet.set_column('E:H', 22); worksheet.set_column('I:I', 10); worksheet.set_column('J:J', 50)
                    
                    topic_sheet = workbook.add_worksheet('主題清單(可擴充)')
                    topic_sheet.write('A1', '🔽 系統預設主題 (往下新增將自動同步)')
                    topic_sheet.set_column('A:A', 50)
                    for i, t in enumerate(["未分類"] + list(topic_mapping_2.keys())):
                        topic_sheet.write(i + 1, 0, t)
                    worksheet.data_validation('C2:C10000', {'validate': 'list', 'source': "='主題清單(可擴充)'!$A$2:$A$200", 'error_type': 'warning'})

                st.success(f"成功解析 {len(qs)} 題！")
                st.download_button("📊 1. 下載待校對 Excel 檔", data=output.getvalue(), file_name=uploaded_word_2.name.replace(".docx", "_校對用.xlsx"), mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True)
                
                img_db_b64 = {k: base64.b64encode(v).decode('utf-8') for k, v in img_db.items()}
                img_json = json.dumps(img_db_b64, ensure_ascii=False)
                st.download_button("🖼️ 2. 下載圖片暫存檔 (image_db.json)", data=img_json, file_name="image_db.json", mime="application/json", use_container_width=True)

with tab3:
    st.info("上傳校對好的 Excel 與圖片暫存檔，打包成最終上線用的 ZIP 題庫包！")
    uploaded_excel = st.file_uploader("1. 上傳校對完的 Excel (.xlsx)", type=["xlsx"])
    uploaded_img_db = st.file_uploader("2. 上傳圖片暫存檔 (image_db.json) (若無可略過)", type=["json"])

    if uploaded_excel and st.button("💾 封裝為最終 ZIP 題庫包", type="primary", use_container_width=True):
        with st.spinner("正在讀取 Excel 並打包實體圖片..."):
            try:
                img_db = {}
                if uploaded_img_db:
                    img_db_b64 = json.load(uploaded_img_db)
                    img_db = {k: base64.b64decode(v) for k, v in img_db_b64.items()}
                
                df = pd.read_excel(uploaded_excel, sheet_name='待校對題庫').fillna("")
                final_questions = []
                for idx, row in df.iterrows():
                    if str(row.get("題目", "")).strip() == "": continue
                    opts = {k: str(row.get(f"選項{k}", "")).strip() for k in ['A', 'B', 'C', 'D'] if str(row.get(f"選項{k}", "")).strip()}
                    q_num = str(row.get("題號", "0")).strip()
                    q_num = int(float(q_num)) if q_num.replace('.', '', 1).isdigit() else 0
                    
                    q = {
                        "question_number": q_num, "answer": str(row.get("正確答案", "")).strip(),
                        "explanation": str(row.get("解析", "")).strip(),
                        "tags": {"年份": str(row.get("年份", "")).strip(), "主題": str(row.get("主題 (下拉選單)", "")).strip()},
                        "question_text": str(row.get("題目", "")).strip(), "options": opts
                    }
                    final_questions.append(q)
                    
                if final_questions:
                    prefix = uploaded_excel.name.replace(".xlsx", "")
                    zip_bytes = create_zip_package(final_questions, img_db, prefix)
                    st.success(f"🎉 封裝成功！共匯入 {len(final_questions)} 題，並打包了 {len(img_db)} 張圖片！")
                    st.download_button("📦 下載 ZIP 題庫包", data=zip_bytes, file_name=f"{prefix}_最終題庫包.zip", mime="application/zip", type="primary", use_container_width=True)
            except Exception as e:
                st.error(f"錯誤：{e}")
